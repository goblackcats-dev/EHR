#!/usr/bin/env python3
"""Builds the printable ADN Med-Surg II clinical care-plan case study (Word) from a content file.

Usage:
  python3 build_casestudy.py content.json patient.json outdir

content.json is written for ONE patient (see BLUEPRINT.md). patient.json is the NursingSim export for that patient (used for the
shift-update page and the vital-sign key). Writes two files into outdir:
  <caseCode>_STUDENT.docx   packet with writing boxes and no answers
  <caseCode>_KEY.docx       faculty answer key, scoring rules and score sheet

Do not publish the key or the content file (keep them out of the public repo / website).
"""
import json, os, sys
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT = 'Times New Roman'
BOX = '☐'


# ---------------------------------------------------------------- low-level helpers
def set_cell_border(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn('w:tcBorders'))
    if b is None:
        b = OxmlElement('w:tcBorders'); tcPr.append(b)
    for edge in ('top', 'left', 'bottom', 'right'):
        v = kw.get(edge)
        if v is None: continue
        el = b.find(qn('w:' + edge))
        if el is None:
            el = OxmlElement('w:' + edge); b.append(el)
        el.set(qn('w:val'), v.get('val', 'single')); el.set(qn('w:sz'), str(v.get('sz', 8))); el.set(qn('w:color'), v.get('color', '000000')); el.set(qn('w:space'), '0')


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), fill); tcPr.append(s)


def set_col_widths(table, widths):
    table.autofit = False
    tblPr = table._tbl.tblPr
    if tblPr.find(qn('w:tblLayout')) is None:
        lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    for i, w in enumerate(widths):
        if i < len(table.columns): table.columns[i].width = Inches(w)
    for row in table.rows:
        for i, w in enumerate(widths):
            if i < len(row.cells): row.cells[i].width = Inches(w)


def row_height(row, inches, exact=False):
    row.height = Inches(inches)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY if exact else WD_ROW_HEIGHT_RULE.AT_LEAST


def no_split(row):
    trPr = row._tr.get_or_add_trPr(); el = OxmlElement('w:cantSplit'); trPr.append(el)


def run(p, text, bold=False, italic=False, size=None, color=None, font=None):
    r = p.add_run(text); r.bold = bold; r.italic = italic
    r.font.name = font or FONT; r._element.rPr.rFonts.set(qn('w:eastAsia'), font or FONT)
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor.from_string(color)
    return r


def para(doc_or_cell, text='', bold=False, italic=False, size=11.5, color=None, align=None, before=0, after=4, keep=False):
    p = doc_or_cell.add_paragraph()
    pf = p.paragraph_format; pf.space_before = Pt(before); pf.space_after = Pt(after)
    if keep: pf.keep_with_next = True
    if align == 'center': p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if text: run(p, text, bold=bold, italic=italic, size=size, color=color)
    return p


def cell_text(cell, text, bold=False, italic=False, size=10.5, color=None, align=None):
    cell.text = ''
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
    if align == 'center': p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, text, bold=bold, italic=italic, size=size, color=color)
    return p


def add_field(paragraph, code, size=9):
    r = paragraph.add_run(); r.font.size = Pt(size); r.font.name = FONT
    for t, txt in (('begin', None), ('instr', code), ('sep', None), ('text', '1'), ('end', None)):
        if t == 'instr':
            el = OxmlElement('w:instrText'); el.set(qn('xml:space'), 'preserve'); el.text = f' {code} '; r._r.append(el)
        elif t == 'text':
            el = OxmlElement('w:t'); el.text = txt; r._r.append(el)
        else:
            el = OxmlElement('w:fldChar'); el.set(qn('w:fldCharType'), {'begin': 'begin', 'sep': 'separate', 'end': 'end'}[t]); r._r.append(el)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


THICK = {'val': 'single', 'sz': 10, 'color': '000000'}
THIN = {'val': 'single', 'sz': 4, 'color': '808080'}
NONE = {'val': 'nil', 'sz': 0, 'color': 'FFFFFF'}


def new_doc(meta, key=False):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.left_margin = s.right_margin = Inches(0.75); s.top_margin = Inches(0.8); s.bottom_margin = Inches(0.7)
    st = d.styles['Normal']; st.font.name = FONT; st.font.size = Pt(11.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    # header: student name / date / case / page numbers on EVERY page so a photo of any page can be matched to its student
    hp = s.header.paragraphs[0]; hp.text = ''
    if key:
        run(hp, f"FACULTY KEY AND SCORING GUIDE  |  {meta['caseCode']}  |  Do not give to students", bold=True, size=9, color='B00000')
    else:
        run(hp, 'Student Name: ______________________________   Date: ____________   ', size=9.5, bold=True)
        run(hp, f"Case {meta['caseCode']}   Page ", size=9.5); add_field(hp, 'PAGE', 9.5); run(hp, ' of ', size=9.5); add_field(hp, 'NUMPAGES', 9.5)
    fp = s.footer.paragraphs[0]; fp.text = ''
    run(fp, meta.get('footer', 'Write one answer per numbered box, inside the box, in dark pen, neatly. Open book, individual work.') if not key else 'Score only what is written. Flag anything you cannot read for faculty review.', size=8.5, italic=True, color='555555')
    return d


# ---------------------------------------------------------------- item renderers (student side)
def item_head(d, it, part):
    p = para(d, before=8, after=2, keep=True)
    run(p, f"{it['id']}  ", bold=True, size=12.5)
    run(p, it.get('title') or '', bold=True, size=12.5)
    run(p, f"   ({fmt(it['points'])} pt{'s' if it['points'] != 1 else ''})", size=10.5, italic=True, color='444444')
    if it.get('prompt'):
        para(d, it['prompt'], size=11.5, after=2, keep=True)
    if it.get('where'):
        para(d, 'Where to look in the EHR: ' + it['where'], italic=True, size=9.5, color='555555', after=3, keep=True)


def fmt(x):
    return str(int(x)) if float(x).is_integer() else str(x)


def ruled_box(d, lines, width=7.0, line_h=0.3, label=None):
    t = d.add_table(rows=lines, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_col_widths(t, [width])
    for i, r in enumerate(t.rows):
        row_height(r, line_h, True); no_split(r)
        c = r.cells[0]
        set_cell_border(c, left=THICK, right=THICK, top=THICK if i == 0 else THIN, bottom=THICK if i == lines - 1 else THIN)
        c.paragraphs[0].paragraph_format.space_after = Pt(0)
        if i < lines - 1: c.paragraphs[0].paragraph_format.keep_with_next = True   # keep the whole writing box on one page
    if label: cell_text(t.rows[0].cells[0], label, size=8, italic=True, color='888888')
    para(d, '', after=2)


def render_fields(d, it):
    fields = it['fields']; per = it.get('perRow', 2)
    rows = (len(fields) + per - 1) // per
    t = d.add_table(rows=rows, cols=per * 2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    lw = it.get('labelW', 1.15); bw = (7.0 - lw * per) / per
    set_col_widths(t, [lw, bw] * per)
    for i, f in enumerate(fields):
        r, c = divmod(i, per)
        lab, blank = t.rows[r].cells[c * 2], t.rows[r].cells[c * 2 + 1]
        cell_text(lab, f['label'] + ':', bold=True, size=10); lab.vertical_alignment = WD_ALIGN_VERTICAL.BOTTOM; blank.vertical_alignment = WD_ALIGN_VERTICAL.BOTTOM
        set_cell_border(blank, bottom=THICK, top=NONE, left=NONE, right=NONE)
        row_height(t.rows[r], it.get('rowH', 0.42), False)
    para(d, '', after=2)


def render_table(d, it):
    cols = it['cols']; widths = it.get('widths') or [7.0 / len(cols)] * len(cols)
    rows = it['rows']; n = rows if isinstance(rows, int) else len(rows)
    t = d.add_table(rows=n + 1, cols=len(cols)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_col_widths(t, widths)
    for j, h in enumerate(cols):
        c = t.rows[0].cells[j]; cell_text(c, h, bold=True, size=9.5); shade(c, 'D9D9D9')
    for i in range(n):
        r = t.rows[i + 1]; row_height(r, it.get('rowH', 0.55)); no_split(r)
        for j in range(len(cols)):
            if not isinstance(rows, int) and j < len(rows[i]) and rows[i][j]:
                cell_text(r.cells[j], str(rows[i][j]), size=10, bold=bool(it.get('boldPre', True)))
    for r in t.rows:
        for c in r.cells: set_cell_border(c, top=THICK, bottom=THICK, left=THICK, right=THICK)
    para(d, '', after=2)


def render_choice(d, it):
    style = it.get('style', 'single')
    if it.get('stem'): para(d, it['stem'], size=11, italic=True, after=2)
    note = {'single': 'Select ONE.', 'multi': 'Select ALL that apply.' if not it.get('select') else f"Select {it['select']}.", 'ordered': 'Number the actions 1 (first) to %d (last).' % len(it['options'])}[style]
    para(d, note, bold=True, size=10, after=2)
    for i, o in enumerate(it['options']):
        p = para(d, after=1)
        p.paragraph_format.left_indent = Inches(0.25)
        if style == 'ordered': run(p, '[ ___ ]  ', bold=True, size=11)
        else: run(p, BOX + '  ', size=13, font='Segoe UI Symbol')
        run(p, o, size=11)
    para(d, '', after=2)


def render_matrix(d, it):
    cols = it['cols']; rows = it['rows']
    t = d.add_table(rows=len(rows) + 1, cols=len(cols) + 1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    first = 7.0 - 1.1 * len(cols); set_col_widths(t, [first] + [1.1] * len(cols))
    cell_text(t.rows[0].cells[0], '', size=9); shade(t.rows[0].cells[0], 'D9D9D9')
    for j, h in enumerate(cols): c = t.rows[0].cells[j + 1]; cell_text(c, h, bold=True, size=9.5, align='center'); shade(c, 'D9D9D9')
    for i, r in enumerate(rows):
        cell_text(t.rows[i + 1].cells[0], r, size=10.5); row_height(t.rows[i + 1], 0.45)
        for j in range(len(cols)): cell_text(t.rows[i + 1].cells[j + 1], BOX, size=14, align='center').runs[0].font.name = 'Segoe UI Symbol'
    for r in t.rows:
        for c in r.cells: set_cell_border(c, top=THICK, bottom=THICK, left=THICK, right=THICK)
    para(d, '', after=2)


def render_bowtie(d, it):
    if it.get('stem'): para(d, it['stem'], size=11, italic=True, after=3)
    t = d.add_table(rows=2, cols=3); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_col_widths(t, [2.45, 2.1, 2.45])
    heads = [f"Actions to Take\n(select {it['pickActions']})", 'Condition the patient is most likely experiencing\n(select 1)', f"Parameters to Monitor\n(select {it['pickParams']})"]
    for j, h in enumerate(heads): c = t.rows[0].cells[j]; cell_text(c, h, bold=True, size=9.5, align='center'); shade(c, 'D9D9D9')
    no_split(t.rows[1]); no_split(t.rows[0])
    for j, lst in enumerate([it['actions'], it['conditions'], it['params']]):
        c = t.rows[1].cells[j]; c.text = ''
        for k, o in enumerate(lst):
            p = c.paragraphs[0] if k == 0 else c.add_paragraph(); p.paragraph_format.space_after = Pt(5)
            run(p, BOX + ' ', size=12, font='Segoe UI Symbol'); run(p, o, size=10)
    for r in t.rows:
        for c in r.cells: set_cell_border(c, top=THICK, bottom=THICK, left=THICK, right=THICK)
    para(d, '', after=2)


def render_concept(d, it):
    labels = it['boxes']
    t = d.add_table(rows=len(labels) * 2 - 1, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_col_widths(t, [1.9, 5.1])
    for i, lab in enumerate(labels):
        r = t.rows[i * 2]; row_height(r, 0.62); cell_text(r.cells[0], lab, bold=True, size=10); shade(r.cells[0], 'EDEDED')
        for c in r.cells: set_cell_border(c, top=THICK, bottom=THICK, left=THICK, right=THICK)
        if i < len(labels) - 1:
            ar = t.rows[i * 2 + 1]; row_height(ar, 0.22, True)
            cell_text(ar.cells[1], '▼', size=10, align='center')
            for c in ar.cells: set_cell_border(c, top=NONE, bottom=NONE, left=NONE, right=NONE)
    para(d, '', after=2)


def shift_update(d, it, patient):
    """Printed 'sealed' page: the later-shift chart notes (from the patient file) plus the instructor's addendum."""
    start = (patient.get('timeline') or {}).get('simulationStart', '')
    bx = d.add_table(rows=1, cols=1); c = bx.rows[0].cells[0]; set_col_widths(bx, [7.0])
    cell_text(c, it.get('banner', 'SEALED PAGE. Do not read until Parts 1 through 5 are finished.'), bold=True, size=12, align='center'); shade(c, 'F2F2F2')
    set_cell_border(c, top=THICK, bottom=THICK, left=THICK, right=THICK); row_height(bx.rows[0], 0.5)
    para(d, it.get('intro', ''), size=11, before=6)
    notes = sorted([n for n in patient.get('notes', []) if n['datetime'] > start], key=lambda n: n['datetime'])
    for n in notes[: it.get('maxNotes', 6)]:
        p = para(d, before=4, after=1); run(p, f"{n['datetime'][5:16].replace('-', '/')}  {n['title']}", bold=True, size=11)
        for line in n['body'].split('\n')[: it.get('maxLines', 9)]:
            line = line.strip()
            if line: para(d, line.replace('**', ''), size=10, after=0)
    for blk in it.get('addendum', []):
        p = para(d, before=6, after=1); run(p, blk['time'] + '  ' + blk['title'], bold=True, size=11, color='1F3864')
        for line in blk['lines']: para(d, '- ' + line, size=10.5, after=0)
    if it.get('table'):
        tb = it['table']; t = d.add_table(rows=len(tb['rows']) + 1, cols=len(tb['cols'])); set_col_widths(t, tb.get('widths') or [7.0 / len(tb['cols'])] * len(tb['cols']))
        for j, h in enumerate(tb['cols']): cc = t.rows[0].cells[j]; cell_text(cc, h, bold=True, size=9.5); shade(cc, 'D9D9D9')
        for i, r in enumerate(tb['rows']):
            for j, v in enumerate(r): cell_text(t.rows[i + 1].cells[j], str(v), size=10)
        for r in t.rows:
            for cc in r.cells: set_cell_border(cc, top=THIN, bottom=THIN, left=THIN, right=THIN)
        para(d, '', after=2)
    if it.get('caption'): para(d, it['caption'], bold=True, size=10.5, before=4)


def vitals_key(patient, times, codes):
    V = {}
    for o in patient['observations']:
        if o['type'] == 'vital': V.setdefault(o['collected'], {})[o['code']] = o['value']
    return [[V.get(t, {}).get(c, '?') for t in times] for c in codes]


def render_item(d, it, patient):
    t = it['type']
    if t == 'text':
        if it.get('title'): para(d, it['title'], bold=True, size=12.5, before=8, after=2, keep=True)
        for ln in it['body'] if isinstance(it['body'], list) else [it['body']]: para(d, ln, size=11, after=3)
        return
    if t == 'sealed': page_break(d); shift_update(d, it, patient); return
    item_head(d, it, None)
    if t == 'fields': render_fields(d, it)
    elif t == 'box': ruled_box(d, it.get('lines', 4), it.get('width', 7.0))
    elif t == 'boxes':
        for sub in it['subs']:
            para(d, sub['label'], bold=True, size=10.5, after=1, keep=True); ruled_box(d, sub.get('lines', 2))
    elif t == 'table': render_table(d, it)
    elif t == 'vitals':
        it2 = dict(it); it2['cols'] = ['Vital sign'] + it['times_label']; it2['rows'] = [[v] for v in it['rowsLabels']]
        it2['widths'] = [1.6] + [(7.0 - 1.6) / len(it['times'])] * len(it['times']); it2['rowH'] = 0.42; render_table(d, it2)
        ruled_box(d, 1, 7.0, 0.38, 'Trend in one sentence:')
    elif t == 'choice': render_choice(d, it)
    elif t == 'matrix': render_matrix(d, it)
    elif t == 'bowtie': render_bowtie(d, it)
    elif t == 'concept': render_concept(d, it)
    elif t == 'highlight':
        para(d, it.get('how', 'Draw a line under each statement that needs follow-up now. Then write the item numbers in the box below.'), size=10, bold=True, after=3)
        for i, s in enumerate(it['statements'], 1): p = para(d, after=3); p.paragraph_format.left_indent = Inches(0.25); run(p, f'({i}) ', bold=True, size=11); run(p, s, size=11)
        ruled_box(d, 1, 7.0, 0.35, 'Item numbers that need follow-up:')
    else: raise SystemExit(f"unknown item type {t}")


# ---------------------------------------------------------------- key side
def key_item(d, it, patient):
    t = it['type']
    if t in ('text', 'sealed'): return
    k = it.get('key', {})
    p = para(d, before=7, after=1, keep=True)
    run(p, f"{it['id']}  {it.get('title') or ''}", bold=True, size=11.5); run(p, f"   [{fmt(it['points'])} pts]", italic=True, size=10, color='444444')
    if it.get('prompt'): para(d, it['prompt'], italic=True, size=9.5, color='555555', after=2)
    if t == 'fields':
        tb = d.add_table(rows=1, cols=3); set_col_widths(tb, [2.0, 4.2, 0.8])
        for j, h in enumerate(['Field', 'Expected answer', 'Pts']): c = tb.rows[0].cells[j]; cell_text(c, h, bold=True, size=9); shade(c, 'D9D9D9')
        for f in it['fields']:
            r = tb.add_row(); cell_text(r.cells[0], f['label'], size=9.5); cell_text(r.cells[1], f['answer'], size=9.5); cell_text(r.cells[2], fmt(f['pts']), size=9.5)
        set_col_widths(tb, [2.0, 4.2, 0.8])
        for r in tb.rows:
            for c in r.cells: set_cell_border(c, top=THIN, bottom=THIN, left=THIN, right=THIN)
    elif t == 'vitals':
        vk = vitals_key(patient, it['times'], it['codes'])
        tb = d.add_table(rows=len(vk) + 1, cols=len(it['times']) + 1)
        cell_text(tb.rows[0].cells[0], 'Vital', bold=True, size=9); shade(tb.rows[0].cells[0], 'D9D9D9')
        for j, h in enumerate(it['times_label']): c = tb.rows[0].cells[j + 1]; cell_text(c, h, bold=True, size=9); shade(c, 'D9D9D9')
        for i, row in enumerate(vk):
            cell_text(tb.rows[i + 1].cells[0], it['rowsLabels'][i], size=9.5)
            for j, v in enumerate(row): cell_text(tb.rows[i + 1].cells[j + 1], v, size=9.5)
        for r in tb.rows:
            for c in r.cells: set_cell_border(c, top=THIN, bottom=THIN, left=THIN, right=THIN)
        para(d, k.get('rules', 'Score = points x (cells correct / total cells), rounded to the nearest 0.5. Accept the stated unit or no unit; allow +/- 2 for rates and +/- 0.2 F.'), size=9.5, italic=True, after=2)
    elif t == 'choice':
        ans = k.get('answer', [])
        para(d, 'Correct: ' + '; '.join(ans if isinstance(ans, list) else [ans]), size=10, bold=True, after=1)
        para(d, k.get('rules', ''), size=9.5, italic=True, after=1)
    elif t == 'matrix':
        for r, a in zip(it['rows'], k['answer']): para(d, f'{r}  ->  {a}', size=10, after=0)
        para(d, k.get('rules', 'Each row is worth equal points; no partial credit within a row.'), size=9.5, italic=True, after=1)
    elif t == 'bowtie':
        para(d, 'Condition: ' + k['condition'], size=10, bold=True, after=0); para(d, 'Actions: ' + '; '.join(k['actions']), size=10, after=0); para(d, 'Parameters: ' + '; '.join(k['params']), size=10, after=1)
        para(d, k.get('rules', 'Condition 1 pt (all or nothing). Actions and parameters: 0.5 pt per correct selection; subtract 0.25 for each incorrect selection in that column (not below 0).'), size=9.5, italic=True, after=1)
    elif t == 'highlight':
        para(d, 'Statements that need follow-up: ' + ', '.join(str(x) for x in k['answer']), size=10, bold=True, after=0); para(d, k.get('rules', 'Equal points per correct statement; subtract the same for each wrong statement chosen (not below 0).'), size=9.5, italic=True, after=1)
    else:
        if k.get('answer'): para(d, 'Model answer: ' + k['answer'], size=10, after=2)
    els = k.get('elements')
    if els:
        tb = d.add_table(rows=1, cols=3); set_col_widths(tb, [5.4, 1.1, 0.5])
        for j, h in enumerate(['Required element (accept equivalent wording)', 'Where / who', 'Pts']): c = tb.rows[0].cells[j]; cell_text(c, h, bold=True, size=9); shade(c, 'D9D9D9')
        for e in els:
            r = tb.add_row(); cell_text(r.cells[0], e['t'], size=9.5); cell_text(r.cells[1], e.get('src', ''), size=8.5); cell_text(r.cells[2], fmt(e['pts']), size=9.5)
        set_col_widths(tb, [5.4, 1.1, 0.5])
        for r in tb.rows:
            for c in r.cells: set_cell_border(c, top=THIN, bottom=THIN, left=THIN, right=THIN)
        tot = sum(e['pts'] for e in els)
        if tot > 0 and not k.get('bank') and abs(tot - it['points']) > 1e-6: print(f"WARNING {it['id']}: elements add to {tot}, item is {it['points']}")
    for lab in ('rules', 'errors'):
        if k.get(lab) and t not in ('choice', 'matrix', 'bowtie', 'highlight', 'vitals'): para(d, ('Scoring: ' if lab == 'rules' else 'Common errors: ') + k[lab], size=9.5, italic=True, after=1)
    if k.get('errors') and t in ('choice', 'matrix', 'bowtie', 'highlight', 'vitals'): para(d, 'Common errors: ' + k['errors'], size=9.5, italic=True, after=1)


def score_sheet(d, content):
    page_break(d)
    para(d, 'SCORE SHEET', bold=True, size=15, after=2)
    para(d, f"Student: ______________________   Date graded: ____________   Passing score: {content['meta']['passPct']} of 100", size=11, after=4)
    rows = []
    for part in content['parts']:
        for it in part['items']:
            if it['type'] in ('text', 'sealed'): continue
            rows.append((part, it))
    tb = d.add_table(rows=1, cols=4); set_col_widths(tb, [0.6, 4.4, 0.8, 1.2])
    for j, h in enumerate(['Item', 'Description', 'Max', 'Earned']): c = tb.rows[0].cells[j]; cell_text(c, h, bold=True, size=9.5); shade(c, 'D9D9D9')
    last = None; tot = 0
    for part, it in rows:
        if part['id'] != last:
            r = tb.add_row(); cell_text(r.cells[0], f"Part {part['id']}", bold=True, size=9.5); cell_text(r.cells[1], part['title'], bold=True, size=9.5)
            cell_text(r.cells[2], fmt(sum(i['points'] for i in part['items'] if i['type'] not in ('text', 'sealed'))), bold=True, size=9.5); cell_text(r.cells[3], '', size=9.5)
            for c in r.cells: shade(c, 'F2F2F2')
            last = part['id']
        r = tb.add_row(); cell_text(r.cells[0], it['id'], size=9); cell_text(r.cells[1], (it.get('title') or it.get('prompt') or '')[:70], size=9); cell_text(r.cells[2], fmt(it['points']), size=9); cell_text(r.cells[3], '', size=9)
        tot += it['points']
    r = tb.add_row(); cell_text(r.cells[1], 'TOTAL', bold=True, size=10.5); cell_text(r.cells[2], fmt(tot), bold=True, size=10.5); cell_text(r.cells[3], '', size=10)
    set_col_widths(tb, [0.6, 4.4, 0.8, 1.2])
    for r in tb.rows:
        for c in r.cells: set_cell_border(c, top=THIN, bottom=THIN, left=THIN, right=THIN)
    para(d, f"\nResult:  [  ] PASS (78 or more)     [  ] DOES NOT PASS (below 78)     Items needing faculty review (illegible or unclear): ______________________", size=10.5, before=6)
    if abs(tot - 100) > 1e-6: print(f"WARNING total points = {tot}, expected 100")


# ---------------------------------------------------------------- documents
def build(content, patient, outdir):
    meta = content['meta']; os.makedirs(outdir, exist_ok=True)
    # ---------------- student packet
    d = new_doc(meta)
    para(d, meta['title'], bold=True, size=18, align='center', after=2)
    para(d, f"{meta['course']}   |   Individual take-home clinical care-plan case study   |   Open book", size=11, align='center', after=8)
    for blk in content.get('studentIntro', []):
        if blk.get('h'): para(d, blk['h'], bold=True, size=12.5, before=6, after=2)
        for ln in blk.get('lines', []): para(d, ln, size=11, after=3)
    # time budget table
    tb = d.add_table(rows=1, cols=3); set_col_widths(tb, [0.8, 4.9, 1.3])
    for j, h in enumerate(['Part', 'Nursing process step', 'Estimated time']): c = tb.rows[0].cells[j]; cell_text(c, h, bold=True, size=10); shade(c, 'D9D9D9')
    total = 0
    for part in content['parts']:
        r = tb.add_row(); cell_text(r.cells[0], part['id'], size=10); cell_text(r.cells[1], f"{part['title']}   ({part['step']})", size=10); cell_text(r.cells[2], f"{part['minutes']} min", size=10); total += part['minutes']
    r = tb.add_row(); cell_text(r.cells[1], 'Estimated total', bold=True, size=10); cell_text(r.cells[2], f"{total // 60} h {total % 60} min", bold=True, size=10)
    set_col_widths(tb, [0.8, 4.9, 1.3])
    for r in tb.rows:
        for c in r.cells: set_cell_border(c, top=THIN, bottom=THIN, left=THIN, right=THIN)
    para(d, '', after=4)
    for blk in content.get('studentRules', []):
        if blk.get('h'): para(d, blk['h'], bold=True, size=12.5, before=6, after=2)
        for ln in blk.get('lines', []): para(d, '- ' + ln, size=11, after=2)
    para(d, 'Honor statement: I completed this work myself, using only my textbooks, notes, drug references and the patient chart.', italic=True, size=10.5, before=8)
    para(d, 'Signature: ______________________________   Date: ______________', size=11, before=6)
    for part in content['parts']:
        items = part['items']
        page_break(d)
        if items and items[0]['type'] == 'sealed':   # the sealed page is its own sheet, before the Part header, so it can be handed out separately
            shift_update(d, items[0], patient); page_break(d); items = items[1:]
        para(d, f"PART {part['id']}: {part['title'].upper()}", bold=True, size=15, after=0)
        para(d, f"Nursing process step: {part['step']}   |   About {part['minutes']} minutes   |   {fmt(sum(i['points'] for i in part['items'] if i['type'] not in ('text', 'sealed')))} points", italic=True, size=10, color='444444', after=4)
        if part.get('intro'): para(d, part['intro'], size=11, after=4)
        for it in items: render_item(d, it, patient)
    d.save(os.path.join(outdir, f"{meta['caseCode']}_STUDENT.docx"))

    # ---------------- faculty key
    k = new_doc(meta, key=True)
    para(k, meta['title'] + ' - FACULTY KEY', bold=True, size=16, after=2)
    para(k, f"Passing score: {meta['passPct']} of 100 points.", bold=True, size=11, after=4)
    for blk in content.get('keyIntro', []):
        if blk.get('h'): para(k, blk['h'], bold=True, size=12, before=6, after=2)
        for ln in blk.get('lines', []): para(k, '- ' + ln, size=10.5, after=2)
    for part in content['parts']:
        page_break(k)
        para(k, f"PART {part['id']}: {part['title']}  ({fmt(sum(i['points'] for i in part['items'] if i['type'] not in ('text', 'sealed')))} pts)", bold=True, size=14, after=3)
        if part.get('keyNote'): para(k, part['keyNote'], italic=True, size=10, after=3)
        for it in part['items']: key_item(k, it, patient)
    score_sheet(k, content)
    k.save(os.path.join(outdir, f"{meta['caseCode']}_KEY.docx"))
    print('wrote', meta['caseCode'] + '_STUDENT.docx', 'and', meta['caseCode'] + '_KEY.docx', 'to', outdir)


if __name__ == '__main__':
    if len(sys.argv) != 4: raise SystemExit(__doc__)
    cj = json.load(open(sys.argv[1])); pj = json.load(open(sys.argv[2]))
    if 'canonical' in pj: pj = pj['canonical']
    build(cj, pj, sys.argv[3])
